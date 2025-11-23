#!/usr/bin/env python3
"""
==================================================================================
SCRIPT: Concatenar archivos Telcel con manejo de fechas Excel serial
Versión en Python equivalente al script R
==================================================================================
"""

import os
import re
import glob
from datetime import datetime, timedelta
from pathlib import Path
from typing import List, Dict, Optional, Tuple

import pandas as pd
import numpy as np
from openpyxl import load_workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter


class TelcelProcessor:
    """Procesador de archivos Telcel con optimización y manejo de encoding"""

    def __init__(self, working_dir: str):
        """
        Inicializar procesador

        Args:
            working_dir: Directorio de trabajo con archivos Telcel
        """
        self.working_dir = Path(working_dir)
        self.encoding_corrections = {
            # Minúsculas
            "Ã¡": "á", "Ã©": "é", "Ã.": "í", "Ã³": "ó", "Ãº": "ú", "Ã±": "ñ",
            # Mayúsculas
            "Ã\u0081": "Á", "Ã‰": "É", "Ã\u008D": "Í", "Ã": "Ó", "Ãš": "Ú", "Ã'": "Ñ",
            # Doble encoding
            "ÃƒÂ¡": "á", "ÃƒÂ©": "é", "ÃƒÂ.": "í", "ÃƒÂ³": "ó", "ÃƒÂº": "ú", "ÃƒÂ±": "ñ",
            # Palabras completas
            "RegiÃ³n": "Región", "DescripciÃ³n": "Descripción",
            "RegiÃƒÂ³n": "Región", "DescripciÃƒÂ³n": "Descripción",
            # Otros
            "Ã¼": "ü", "Ãœ": "Ü", "Ã§": "ç"
        }

        self.column_mapping = {
            "Region": "Región", "region": "Región", "REGION": "Región",
            "NombreVenta": "Nombre Venta", "nombreventa": "Nombre Venta",
            "Nombre_Venta": "Nombre Venta", "NombredeVenta": "Nombre Venta",
            "WEEK": "Week", "WeekNum": "Week", "weeknum": "Week", "Semana": "Week",
            "Descripcion": "Descripción", "DESCRIPCION": "Descripción"
        }

    def fix_encoding(self, text: str) -> str:
        """
        Corregir encoding de texto

        Args:
            text: Texto a corregir

        Returns:
            Texto con encoding corregido
        """
        if pd.isna(text) or not isinstance(text, str):
            return text

        for pattern, replacement in self.encoding_corrections.items():
            text = text.replace(pattern, replacement)

        return text

    def fix_encoding_series(self, series: pd.Series) -> pd.Series:
        """Aplicar corrección de encoding a una serie de pandas"""
        return series.apply(lambda x: self.fix_encoding(x) if isinstance(x, str) else x)

    def standardize_columns(self, df: pd.DataFrame, verbose: bool = False) -> pd.DataFrame:
        """
        Estandarizar nombres de columnas

        Args:
            df: DataFrame a estandarizar
            verbose: Mostrar cambios

        Returns:
            DataFrame con columnas estandarizadas
        """
        # Primero corregir encoding
        df.columns = [self.fix_encoding(str(col)) for col in df.columns]

        # Luego aplicar mapeo
        new_columns = []
        for col in df.columns:
            if col in self.column_mapping:
                new_col = self.column_mapping[col]
                if verbose:
                    print(f"     Renombrando: '{col}' → '{new_col}'")
                new_columns.append(new_col)
            else:
                new_columns.append(col)

        df.columns = new_columns
        return df

    @staticmethod
    def excel_serial_to_date(serial_num: float) -> pd.Timestamp:
        """
        Convertir número serial de Excel a fecha

        Args:
            serial_num: Número serial de Excel

        Returns:
            Fecha como Timestamp
        """
        # Excel cuenta desde 1900-01-01 (con bug de año bisiesto)
        if serial_num > 59:
            return pd.Timestamp('1899-12-30') + pd.Timedelta(days=serial_num)
        else:
            return pd.Timestamp('1899-12-31') + pd.Timedelta(days=serial_num)

    def process_dates(self, fecha_series: pd.Series) -> pd.Series:
        """
        Procesar columna de fechas con múltiples formatos

        Args:
            fecha_series: Serie con fechas en diferentes formatos

        Returns:
            Serie con fechas estandarizadas
        """
        fechas_convertidas = pd.Series([pd.NaT] * len(fecha_series), dtype='datetime64[ns]')

        for i, fecha_val in enumerate(fecha_series):
            if pd.isna(fecha_val) or fecha_val == "":
                continue

            fecha_str = str(fecha_val)

            # Detectar número serial de Excel (5 dígitos)
            if re.match(r'^\d{5}$', fecha_str):
                try:
                    fechas_convertidas[i] = self.excel_serial_to_date(float(fecha_str))
                except:
                    pass
            else:
                # Intentar diferentes formatos
                for fmt in ['%Y-%m-%d', '%d/%m/%Y', '%d-%m-%Y', '%Y/%m/%d', '%Y-%m-%d %H:%M:%S']:
                    try:
                        fechas_convertidas[i] = pd.to_datetime(fecha_str, format=fmt)
                        break
                    except:
                        continue

        return fechas_convertidas

    def process_week_column(self, week_series: pd.Series, verbose: bool = False) -> pd.Series:
        """
        Procesar columna Week para formato W## consistente

        Args:
            week_series: Serie con valores de semana
            verbose: Mostrar progreso

        Returns:
            Serie con formato W## estandarizado
        """
        week_processed = week_series.astype(str)

        # Convertir valores numéricos a W##
        numeric_mask = week_processed.str.match(r'^\d+$')
        if numeric_mask.any():
            week_processed[numeric_mask] = week_processed[numeric_mask].apply(
                lambda x: f"W{int(x):02d}"
            )
            if verbose:
                print(f"     Convertidos {numeric_mask.sum()} valores numéricos a W##")

        # Convertir W# a W0#
        single_digit_mask = week_processed.str.match(r'^W\d$')
        if single_digit_mask.any():
            week_processed[single_digit_mask] = week_processed[single_digit_mask].apply(
                lambda x: f"W{int(x[1]):02d}"
            )
            if verbose:
                print(f"     Reformateados {single_digit_mask.sum()} valores W# a W0#")

        return week_processed

    def process_data_types(self, df: pd.DataFrame, verbose: bool = False) -> pd.DataFrame:
        """
        Procesar tipos de datos del DataFrame

        Args:
            df: DataFrame a procesar
            verbose: Mostrar progreso

        Returns:
            DataFrame procesado
        """
        # Convertir a string para evitar conflictos en concatenación
        char_cols = ['Material', 'Cantidad', 'Importe', 'Week', 'Venta']
        for col in char_cols:
            if col in df.columns:
                df[col] = df[col].astype(str)

        # Convertir Fecha a string temporalmente
        if 'Fecha' in df.columns:
            df['Fecha'] = df['Fecha'].astype(str)

        # Procesar columna Week
        if 'Week' in df.columns:
            df['Week'] = self.process_week_column(df['Week'], verbose)

        # Aplicar corrección de encoding a columnas de texto
        for col in df.select_dtypes(include=['object']).columns:
            df[col] = self.fix_encoding_series(df[col])

        return df

    def restore_numeric_columns(self, df: pd.DataFrame) -> pd.DataFrame:
        """
        Restaurar tipos numéricos después de combinar

        Args:
            df: DataFrame combinado

        Returns:
            DataFrame con tipos restaurados
        """
        # Convertir columnas numéricas
        numeric_cols = ['Cantidad', 'Importe', 'Material']
        for col in numeric_cols:
            if col in df.columns:
                df[col] = pd.to_numeric(df[col], errors='coerce')

        # Procesar fechas
        if 'Fecha' in df.columns:
            print("   Procesando fechas...")
            df['Fecha'] = self.process_dates(df['Fecha'])

            fechas_validas = df['Fecha'].notna().sum()
            total = len(df['Fecha'])
            pct = (fechas_validas / total * 100) if total > 0 else 0
            print(f"   Fechas convertidas: {fechas_validas} de {total} ({pct:.1f}%)")

        return df

    def get_target_sheet(self, file_path: Path) -> str:
        """
        Determinar qué hoja del Excel usar

        Args:
            file_path: Ruta del archivo Excel

        Returns:
            Nombre de la hoja a usar
        """
        xl_file = pd.ExcelFile(file_path)
        sheet_names = xl_file.sheet_names

        # Buscar hoja con patrón R+números
        r_pattern_sheets = [s for s in sheet_names if re.match(r'^R\d+$', s)]

        return r_pattern_sheets[0] if r_pattern_sheets else sheet_names[0]

    def process_telcel_file(self, file_path: Path, verbose: bool = True) -> Dict:
        """
        Procesar un archivo Telcel

        Args:
            file_path: Ruta del archivo
            verbose: Mostrar progreso

        Returns:
            Diccionario con resultado del procesamiento
        """
        if verbose:
            print(f"\n📂 Procesando: {file_path.name}")

        try:
            # Leer archivo
            target_sheet = self.get_target_sheet(file_path)
            df = pd.read_excel(file_path, sheet_name=target_sheet)

            if verbose:
                print(f"   📄 Hoja: '{target_sheet}' | 📊 {len(df)} filas × {len(df.columns)} columnas")

            # Estandarizar y procesar
            df = self.standardize_columns(df, verbose=verbose)
            df = self.process_data_types(df, verbose=verbose)

            df['source_file'] = file_path.name

            if verbose:
                print("   ✅ Procesamiento completado")

            return {
                'data': df,
                'filename': file_path.name,
                'success': True
            }

        except Exception as e:
            if verbose:
                print(f"   ❌ ERROR: {str(e)}")

            return {
                'data': None,
                'filename': file_path.name,
                'success': False,
                'error': str(e)
            }

    def create_week_analysis(self, df: pd.DataFrame) -> Optional[pd.DataFrame]:
        """
        Crear análisis por semana

        Args:
            df: DataFrame combinado

        Returns:
            DataFrame con análisis o None si faltan columnas
        """
        required_cols = ['Week', 'Fecha', 'Región']
        if not all(col in df.columns for col in required_cols):
            return None

        # Filtrar fechas válidas
        df_valid = df[df['Fecha'].notna()].copy()

        analysis = df_valid.groupby('Week').agg({
            'Fecha': ['min', 'max', 'nunique'],
            'Región': lambda x: ', '.join(sorted(x.unique()))
        }).reset_index()

        analysis.columns = ['Week', 'Fecha_Inicio', 'Fecha_Fin', 'Dias_Unicos', 'Regiones']

        # Añadir total de registros
        total_registros = df.groupby('Week').size().reset_index(name='Total_Registros')
        analysis = analysis.merge(total_registros, on='Week')

        # Fechas presentes
        fechas_presentes = df_valid.groupby('Week')['Fecha'].apply(
            lambda x: ', '.join(sorted(x.dt.strftime('%d/%m/%Y').unique()))
        ).reset_index(name='Fechas_Presentes')
        analysis = analysis.merge(fechas_presentes, on='Week')

        # Rango de fechas
        analysis['Rango_Fechas'] = (
            analysis['Fecha_Inicio'].dt.strftime('%d/%m/%Y') + ' a ' +
            analysis['Fecha_Fin'].dt.strftime('%d/%m/%Y')
        )

        # Reordenar columnas
        analysis = analysis[[
            'Week', 'Rango_Fechas', 'Dias_Unicos',
            'Total_Registros', 'Regiones', 'Fechas_Presentes'
        ]]

        return analysis.sort_values('Week')

    def save_formatted_excel(self, df: pd.DataFrame, filename: str):
        """
        Guardar Excel con formato de fechas

        Args:
            df: DataFrame a guardar
            filename: Nombre del archivo
        """
        # Formatear fechas antes de guardar
        df_to_save = df.copy()
        if 'Fecha' in df_to_save.columns:
            df_to_save['Fecha'] = pd.to_datetime(df_to_save['Fecha']).dt.strftime('%d/%m/%Y')

        # Guardar con openpyxl para mejor control
        df_to_save.to_excel(filename, index=False, engine='openpyxl')

        # Aplicar estilos
        wb = load_workbook(filename)
        ws = wb.active

        # Estilo de encabezado
        header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        header_font = Font(color="FFFFFF", bold=True, size=12)

        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal='center')

        # Ajustar ancho de columnas
        for column in ws.columns:
            max_length = 0
            column_letter = get_column_letter(column[0].column)
            for cell in column:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = min(max_length + 2, 50)
            ws.column_dimensions[column_letter].width = adjusted_width

        wb.save(filename)

    def process_all_files(self) -> Tuple[pd.DataFrame, List[Dict]]:
        """
        Procesar todos los archivos Telcel en el directorio

        Returns:
            Tupla con (DataFrame combinado, lista de resultados)
        """
        print("🚀 PROCESAMIENTO DE ARCHIVOS TELCEL - VERSIÓN PYTHON")
        print("=" * 70)
        print(f"📅 Inicio: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}\n")

        # Buscar archivos
        pattern = self.working_dir / "Telcel_WK*.xlsx"
        files = sorted(glob.glob(str(pattern)))

        print(f"📁 Archivos encontrados: {len(files)}")

        if len(files) == 0:
            raise FileNotFoundError("❌ No se encontraron archivos para procesar")

        # Mostrar archivos
        print("\n📋 Archivos a procesar:")
        for i, file_path in enumerate(files, 1):
            print(f"   {i:2d}. {Path(file_path).name}")

        # Procesar archivos
        print("\n⚙️  PROCESAMIENTO")
        print("=" * 50)

        results = [self.process_telcel_file(Path(f), verbose=True) for f in files]
        successful_data = [r['data'] for r in results if r['success']]

        # Resumen
        successful_count = sum(r['success'] for r in results)
        print(f"\n📊 RESUMEN:")
        print(f"   ✅ Exitosos: {successful_count}")
        print(f"   ❌ Errores: {len(files) - successful_count}")

        # Combinar datos
        if successful_data:
            print("\n🔗 COMBINANDO DATOS")
            print("=" * 40)

            combined = pd.concat(successful_data, ignore_index=True)
            combined = self.restore_numeric_columns(combined)

            # Remover columna auxiliar
            combined_for_save = combined.drop(columns=['source_file'], errors='ignore')

            print(f"\n✅ Datos combinados:")
            print(f"   • Filas: {len(combined):,}")
            print(f"   • Columnas: {len(combined_for_save.columns)}")

            # Distribución por semana
            if 'Week' in combined.columns:
                print("\n📅 Distribución por semana:")
                week_dist = combined['Week'].value_counts().sort_index()
                total = len(combined)
                for week, count in week_dist.items():
                    pct = (count / total * 100) if total > 0 else 0
                    print(f"   {week}: {count:,} ({pct:.1f}%)")

            return combined_for_save, results
        else:
            raise RuntimeError("❌ No se procesaron archivos exitosamente")

    def save_all_outputs(self, df: pd.DataFrame):
        """
        Guardar todos los archivos de salida

        Args:
            df: DataFrame a guardar
        """
        print("\n💾 GUARDANDO ARCHIVOS")
        print("=" * 40)

        # Excel principal
        output_xlsx = self.working_dir / "Telcel_ALL_WK_sep_Reg.xlsx"
        self.save_formatted_excel(df, str(output_xlsx))
        print(f"   ✅ Excel: {output_xlsx.name}")

        # CSV
        output_csv = self.working_dir / "Telcel_ALL_WK_Sep_Reg.csv"
        df_csv = df.copy()
        if 'Fecha' in df_csv.columns:
            df_csv['Fecha'] = pd.to_datetime(df_csv['Fecha']).dt.strftime('%d/%m/%Y')
        df_csv.columns = [self.fix_encoding(str(col)) for col in df_csv.columns]
        df_csv.to_csv(output_csv, index=False, encoding='utf-8-sig')
        print(f"   ✅ CSV: {output_csv.name}")

        # Reporte de análisis
        week_analysis = self.create_week_analysis(df)

        if week_analysis is not None:
            timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
            report_filename = self.working_dir / f"Telcel_Analisis_{timestamp}.xlsx"

            self.save_formatted_excel(week_analysis, str(report_filename))
            print(f"   ✅ Reporte: {report_filename.name}")

        # Resumen final
        print("\n🎯 COMPLETADO")
        print("=" * 40)
        print(f"   • Total registros: {len(df):,}")


def main():
    """Función principal"""
    # Configurar directorio de trabajo
    # Modificar según necesidad
    working_dir = "C:/Users/FROJAS/Desktop/ArchivosTelcelAnalisisSemanas/SemanasTabletas"

    try:
        processor = TelcelProcessor(working_dir)
        combined_df, results = processor.process_all_files()
        processor.save_all_outputs(combined_df)

        print(f"\n✨ Finalizado: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")

    except Exception as e:
        print(f"\n❌ Error crítico: {str(e)}")
        raise


if __name__ == "__main__":
    main()
